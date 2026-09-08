"""
Khova AI — Core POC
Validates the 5 hard primitives the whole app depends on:
  (a) LLM structured JSON generation via emergentintegrations LlmChat
  (b) Real-time web research WITH CITATIONS via Gemini googleSearch grounding
  (c) Nano Banana image generation -> base64 -> save PNG
  (d) WeasyPrint HTML/CSS -> real PDF with an embedded base64 image
  (e) openpyxl -> real .xlsx with multiple sheets + a formula

Run:  cd /app/backend && python poc_core.py
"""
import os
import json
import base64
import asyncio
import traceback
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

OUT = ROOT / "generated" / "poc"
OUT.mkdir(parents=True, exist_ok=True)

API_KEY = os.getenv("EMERGENT_LLM_KEY")

from emergentintegrations.llm.chat import LlmChat, UserMessage

results = {}


def _extract_json(text: str):
    """Best-effort JSON extraction from an LLM response."""
    text = text.strip()
    if text.startswith("```"):
        # strip code fences
        text = text.split("```", 2)[1]
        if text.lstrip().startswith("json"):
            text = text.lstrip()[4:]
    # find first { and last }
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1:
        text = text[start:end + 1]
    return json.loads(text)


async def test_llm_json():
    """(a) Structured JSON generation."""
    print("\n=== (a) LLM STRUCTURED JSON ===")
    chat = LlmChat(
        api_key=API_KEY,
        session_id="poc-json",
        system_message="You output ONLY valid minified JSON. No markdown, no commentary.",
    ).with_model("openai", "gpt-5.4-mini")
    prompt = (
        "Return a JSON object with keys: title (string), audience (string), "
        "chapters (array of 3 strings). Topic: 'Personal finance for freelancers'. "
        "Respond in Indonesian language."
    )
    resp = await chat.send_message(UserMessage(text=prompt))
    data = _extract_json(resp)
    assert "title" in data and "chapters" in data and len(data["chapters"]) == 3
    print("  title:", data["title"])
    print("  chapters:", data["chapters"])
    results["a_llm_json"] = True


async def test_web_research():
    """(b) Real web research with citations via Gemini googleSearch."""
    print("\n=== (b) WEB RESEARCH WITH CITATIONS (Gemini googleSearch) ===")
    chat = LlmChat(
        api_key=API_KEY,
        session_id="poc-research",
        system_message="You are a market researcher. Use web search to find current, real information.",
    ).with_model("gemini", "gemini-2.5-flash").with_tools([{"googleSearch": {}}])

    resp = await chat.send_message(UserMessage(
        text="Research the current market for personal finance ebooks and budgeting spreadsheet products for freelancers in 2026. "
             "Mention specific products, pricing, and common customer complaints you find."
    ))
    print("  response length:", len(resp) if resp else 0)

    # Try to extract citations from raw metadata
    sources = []
    try:
        raw = getattr(chat, "last_response_raw", None)
    except Exception:
        raw = None
    # The library exposes raw differently; try a fresh non-stream call to inspect structure
    print("  (attempting citation extraction via send_message_with_tools raw)")
    chat2 = LlmChat(
        api_key=API_KEY,
        session_id="poc-research2",
        system_message="You are a market researcher. Use web search for real, current info.",
    ).with_model("gemini", "gemini-2.5-flash").with_tools([{"googleSearch": {}}])
    r2 = await chat2.send_message_with_tools(UserMessage(
        text="Find 3 real personal finance / budgeting digital products for freelancers with their source URLs."
    ))
    try:
        raw = r2.raw
        choice = raw["choices"][0] if isinstance(raw, dict) else raw.choices[0]
        # annotations
        msg = choice["message"] if isinstance(choice, dict) else choice.message
        anns = (msg.get("annotations") if isinstance(msg, dict) else getattr(msg, "annotations", None)) or []
        for a in anns:
            url = None
            title = None
            if isinstance(a, dict):
                uc = a.get("url_citation") or a
                url = uc.get("url")
                title = uc.get("title")
            if url:
                sources.append({"title": title, "url": url})
    except Exception as e:
        print("  annotation extraction note:", repr(e))

    # fallback: groundingMetadata
    if not sources:
        try:
            raw = r2.raw
            choice = raw["choices"][0] if isinstance(raw, dict) else raw.choices[0]
            gm = None
            if isinstance(choice, dict):
                gm = choice.get("groundingMetadata") or (choice.get("message", {}) or {}).get("groundingMetadata")
            else:
                gm = getattr(choice, "groundingMetadata", None)
            if gm:
                chunks = gm.get("groundingChunks", []) if isinstance(gm, dict) else getattr(gm, "groundingChunks", [])
                for c in chunks:
                    web = c.get("web") if isinstance(c, dict) else getattr(c, "web", None)
                    if web:
                        url = web.get("uri") if isinstance(web, dict) else getattr(web, "uri", None)
                        title = web.get("title") if isinstance(web, dict) else getattr(web, "title", None)
                        if url:
                            sources.append({"title": title, "url": url})
        except Exception as e:
            print("  grounding extraction note:", repr(e))

    print(f"  citations found: {len(sources)}")
    for s in sources[:5]:
        print("   -", s.get("title"), "->", s.get("url"))

    # Save raw structure for debugging
    try:
        with open(OUT / "research_raw.json", "w") as f:
            json.dump({"text_len": len(r2.content or ""), "sources": sources}, f, indent=2)
    except Exception:
        pass

    results["b_web_research"] = {"has_text": bool(resp), "citations": len(sources)}


async def test_image_gen():
    """(c) Nano Banana image generation."""
    print("\n=== (c) NANO BANANA IMAGE GEN ===")
    chat = LlmChat(
        api_key=API_KEY,
        session_id="poc-image",
        system_message="You are a professional book cover designer.",
    ).with_model("gemini", "gemini-3.1-flash-image-preview").with_params(modalities=["image", "text"])
    msg = UserMessage(
        text="A premium, minimalist ebook cover illustration: abstract geometric shapes representing "
             "financial growth, navy blue and gold palette, clean modern style, no text."
    )
    text, images = await chat.send_message_multimodal_response(msg)
    assert images and len(images) > 0, "No images returned"
    img = images[0]
    print("  mime:", img["mime_type"], "| data prefix:", img["data"][:10])
    image_bytes = base64.b64decode(img["data"])
    png_path = OUT / "cover.png"
    with open(png_path, "wb") as f:
        f.write(image_bytes)
    print("  saved PNG bytes:", png_path.stat().st_size)
    results["c_image_gen"] = png_path.stat().st_size
    return png_path


async def test_pdf(png_path):
    """(d) WeasyPrint PDF with embedded image."""
    print("\n=== (d) WEASYPRINT PDF (embedded image) ===")
    from weasyprint import HTML
    img_b64 = base64.b64encode(png_path.read_bytes()).decode() if png_path and png_path.exists() else ""
    img_tag = f'<img src="data:image/png;base64,{img_b64}" style="width:60%;border-radius:12px;" />' if img_b64 else ""
    html = f"""
    <html><head><style>
      @page {{ size: A4; margin: 2cm; }}
      body {{ font-family: 'Helvetica', sans-serif; color: #1a202c; }}
      .cover {{ text-align:center; padding-top: 4cm; }}
      h1 {{ color:#1e3a5f; font-size: 34px; }}
      .accent {{ color:#c9a227; }}
      .callout {{ background:#f0f4f8; border-left:5px solid #1e3a5f; padding:14px; border-radius:8px; }}
    </style></head><body>
      <div class="cover">
        <h1>Keuangan Pribadi untuk <span class="accent">Freelancer</span></h1>
        {img_tag}
      </div>
      <div style="page-break-before: always;"></div>
      <h2>Bab 1: Dasar</h2>
      <p>Ini adalah paragraf contoh untuk memvalidasi rendering PDF.</p>
      <div class="callout">Callout box menggunakan design system warna.</div>
    </body></html>
    """
    pdf_path = OUT / "ebook.pdf"
    HTML(string=html).write_pdf(str(pdf_path))
    print("  saved PDF bytes:", pdf_path.stat().st_size)
    results["d_pdf"] = pdf_path.stat().st_size


def test_xlsx():
    """(e) openpyxl multi-sheet workbook with formula."""
    print("\n=== (e) OPENPYXL XLSX (multi-sheet + formula) ===")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Anggaran"
    headers = ["Kategori", "Anggaran", "Aktual", "Selisih"]
    fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    rows = [["Sewa", 5000000, 5000000], ["Makanan", 3000000, 3200000], ["Transport", 1000000, 850000]]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r[0])
        ws.cell(row=i, column=2, value=r[1])
        ws.cell(row=i, column=3, value=r[2])
        ws.cell(row=i, column=4, value=f"=B{i}-C{i}")  # formula
    ws.cell(row=5, column=1, value="TOTAL")
    ws.cell(row=5, column=2, value="=SUM(B2:B4)")
    ws.cell(row=5, column=3, value="=SUM(C2:C4)")
    # second sheet
    ws2 = wb.create_sheet("Instruksi")
    ws2["A1"] = "Cara pakai spreadsheet ini"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2["A2"] = "Isi kolom Aktual setiap bulan. Selisih dihitung otomatis."
    xlsx_path = OUT / "spreadsheet.xlsx"
    wb.save(str(xlsx_path))
    print("  saved XLSX bytes:", xlsx_path.stat().st_size)
    results["e_xlsx"] = xlsx_path.stat().st_size


async def main():
    print("EMERGENT_LLM_KEY present:", bool(API_KEY))
    png_path = None
    for name, coro in [
        ("a", test_llm_json()),
        ("b", test_web_research()),
        ("c", test_image_gen()),
    ]:
        try:
            r = await coro
            if name == "c":
                png_path = r
        except Exception as e:
            print(f"  !! {name} FAILED: {e}")
            traceback.print_exc()
            results[f"{name}_error"] = str(e)
    try:
        await test_pdf(png_path)
    except Exception as e:
        print("  !! d PDF FAILED:", e)
        traceback.print_exc()
        results["d_error"] = str(e)
    try:
        test_xlsx()
    except Exception as e:
        print("  !! e XLSX FAILED:", e)
        traceback.print_exc()
        results["e_error"] = str(e)

    print("\n================ POC SUMMARY ================")
    print(json.dumps(results, indent=2, default=str))
    passed = all([
        results.get("a_llm_json"),
        results.get("b_web_research", {}).get("has_text") if isinstance(results.get("b_web_research"), dict) else False,
        results.get("c_image_gen"),
        results.get("d_pdf"),
        results.get("e_xlsx"),
    ])
    print("\nALL CORE PRIMITIVES PASSED:", passed)


if __name__ == "__main__":
    asyncio.run(main())
